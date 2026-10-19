"""Export engine: intent detection, every format, and awkward real-world data."""
import csv
import datetime
import decimal
import io
import json

import pytest

from reporting.export_engine import (
    ExportFormat, ExportRequest, detect_export_intent, generate_export,
)


@pytest.mark.parametrize("q,fmt", [
    ("صدّر النتائج pdf", ExportFormat.PDF),
    ("export to excel", ExportFormat.EXCEL),
    ("حول الجدول لـ اكسل", ExportFormat.EXCEL),
    ("export as csv", ExportFormat.CSV),
    ("export this as json", ExportFormat.JSON),
    ("صدر التقرير", ExportFormat.EXCEL),
    ("download data", ExportFormat.EXCEL),
])
def test_export_intent_detected(q, fmt):
    assert detect_export_intent(q) == fmt


@pytest.mark.parametrize("q", ["كام رصيد العميل", "show me sales", "", "ما هو ال pdf"])
def test_non_export_questions_are_ignored(q):
    assert detect_export_intent(q) is None


def _req(fmt, data=None, **kw):
    return ExportRequest(format=fmt, title=kw.pop("title", "تقرير المبيعات"),
                         data=[{"الاسم": "أحمد", "الرصيد": 10.5}] if data is None else data, **kw)


def test_csv_has_bom_header_and_rows():
    content, name, media = generate_export(_req(ExportFormat.CSV, filters={"الشركة": "1"}))
    text = content.decode("utf-8")
    assert text.startswith("﻿") and "الفلاتر: الشركة=1" in text
    body = [l for l in text.splitlines() if l and not l.startswith(("﻿#", "#"))]
    assert list(csv.reader(body))[1] == ["أحمد", "10.5"]
    assert name.endswith(".csv") and media.startswith("text/csv")


def test_csv_empty_data_message():
    content, _, _ = generate_export(_req(ExportFormat.CSV, data=[]))
    assert "لا توجد بيانات" in content.decode("utf-8")


def test_csv_rows_with_different_keys_do_not_crash():
    rows = [{"a": 1}, {"a": 2, "b": 3}]
    content, _, _ = generate_export(_req(ExportFormat.CSV, data=rows))
    assert "a" in content.decode("utf-8")


@pytest.mark.parametrize("payload", ["=HYPERLINK(\"http://x\")", "+1+1", "-2+3", "@SUM(A1)"])
def test_csv_neutralises_spreadsheet_formulas(payload):
    content, _, _ = generate_export(_req(ExportFormat.CSV, data=[{"note": payload}]))
    cell = list(csv.reader(io.StringIO(content.decode("utf-8").split("\n\n", 1)[1])))[1][0]
    assert not cell.startswith(("=", "+", "-", "@"))


def test_excel_roundtrip_and_odd_values():
    import openpyxl
    rows = [{"d": datetime.date(2024, 1, 5), "n": decimal.Decimal("12.50"), "none": None,
             "bad": "a\x00b\x07c", "lst": [1, 2], "ok": 3}]
    content, name, media = generate_export(_req(ExportFormat.EXCEL, data=rows, kpis=[{"label": "k", "value": 1}],
                                               summary="ملخص"))
    ws = openpyxl.load_workbook(io.BytesIO(content)).active
    values = [c.value for r in ws.iter_rows() for c in r if c.value is not None]
    assert "ملخص" in values and 3 in values
    assert name.endswith(".xlsx") and "spreadsheetml" in media


def test_excel_neutralises_formulas():
    import openpyxl
    content, _, _ = generate_export(_req(ExportFormat.EXCEL, data=[{"note": "=1+1"}]))
    ws = openpyxl.load_workbook(io.BytesIO(content)).active
    cells = [c for r in ws.iter_rows() for c in r if c.value is not None and "1+1" in str(c.value)]
    assert cells and all(c.data_type != "f" for c in cells)


def test_excel_sheet_title_is_truncated():
    import openpyxl
    content, _, _ = generate_export(_req(ExportFormat.EXCEL, title="ع" * 80))
    assert len(openpyxl.load_workbook(io.BytesIO(content)).active.title) <= 31


def test_json_serialises_dates_and_decimals():
    rows = [{"d": datetime.date(2024, 1, 5), "n": decimal.Decimal("1.5")}]
    content, _, media = generate_export(_req(ExportFormat.JSON, data=rows))
    parsed = json.loads(content)
    assert parsed["row_count"] == 1 and parsed["data"][0]["d"] == "2024-01-05"
    assert media.startswith("application/json")


def test_pdf_html_escapes_values():
    rows = [{"name": "<script>alert(1)</script>"}]
    content, _, media = generate_export(_req(ExportFormat.PDF, data=rows, title="<b>t</b>", summary="<i>s</i>",
                                             filters={"<f>": "<v>"}))
    html = content.decode("utf-8")
    assert "<script>" not in html and "<i>s</i>" not in html and "<b>t</b>" not in html
    assert media.startswith("text/html")


def test_filename_is_sanitised():
    _, name, _ = generate_export(_req(ExportFormat.JSON, title='../../etc/passwd:*?"'))
    assert "/" not in name and ".." not in name and ":" not in name


def test_unknown_format_raises():
    with pytest.raises(ValueError):
        generate_export(_req("docx"))
